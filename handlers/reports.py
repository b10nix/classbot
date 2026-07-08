import os
import io
import csv
import re
import logging
import openpyxl
from openpyxl.styles import PatternFill, Font
from telegram import Update
from telegram.ext import ContextTypes
from telegram.constants import ParseMode
from database.db import get_db

# Log ထုတ်ရန် သတ်မှတ်ခြင်း
logger = logging.getLogger(__name__)

def is_super_admin(user_id: int) -> bool:
    owner_id = int(os.getenv("OWNER_ID", "0"))
    if user_id == owner_id: return True
    sa_str = os.getenv("SUPER_ADMINS", "")
    if sa_str:
        try:
            return user_id in [int(x.strip()) for x in sa_str.split(",") if x.strip()]
        except ValueError: pass
    return False

# ================= 📋 ကျောင်းသားစာရင်းကြည့်ရန် (/list) =================
async def list_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_super_admin(update.effective_user.id): return
    
    db = get_db()
    students = await db.student_info.find({}).to_list(length=None)
    if not students: return await update.message.reply_text("စာရင်းမရှိသေးပါ။")

    registered_nums = {int(re.search(r'(\d+)$', s.get("roll_no", "")).group(1)) for s in students if re.search(r'(\d+)$', s.get("roll_no", ""))}
    max_roll = max(registered_nums) if registered_nums else 59
    if max_roll < 59: max_roll = 59 

    missing_rolls = [i for i in range(1, max_roll + 1) if i not in registered_nums]
    missing_str = ", ".join([f"{i:02d}" for i in missing_rolls])

    msg = (
        "📊 <b>ကျောင်းသား စာရင်းသွင်းမှု အခြေအနေ</b>\n━━━━━━━━━━━━━━━━━━━━━\n\n"
        f"🔹 <b>စာရင်းသွင်းပြီးသူ:</b> {len(students)} ယောက်\n"
        f"🔹 <b>စာရင်းသွင်းရန် ကျန်ရှိသူ:</b> {len(missing_rolls)} ယောက်\n\n"
        "⏳ <b>စာရင်းသွင်းရန် ကျန်သေးသည့် ခုံနံပါတ်များ:</b>\n"
        f"<code>{missing_str if missing_str else 'အားလုံးသွင်းပြီးပါပြီ 🎉'}</code>"
    )
    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)

# ================= 🏠 အဆောင်နေ ကျောင်းသားများ (/hostel) =================
async def hostel_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_super_admin(update.effective_user.id): return
    
    db = get_db()
    students = await db.student_info.find({"hostel": {"$in": ["Yes", "yes", True]}}).sort("roll_no", 1).to_list(length=None)
    if not students: return await update.message.reply_text("အဆောင်နေ ကျောင်းသား မရှိသေးပါ။")
    
    msg = "🏠 <b>အဆောင်နေ ကျောင်းသားများ</b>\n━━━━━━━━━━━━━━\n"
    for s in students: msg += f"🔹 {s.get('roll_no')} - {s.get('name')}\n"
    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)

# ================= 🏦 ရန်ပုံငွေ အခြေအနေချုပ် (/vault) =================
async def vault_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_super_admin(update.effective_user.id): return
    
    db = get_db()
    funds = await db.funds.find({}).to_list(length=None)
    if not funds: return await update.message.reply_text("ရန်ပုံငွေ ခေါင်းစဉ်များ မရှိသေးပါ။")
    
    students = await db.student_info.find({}).to_list(length=None)
    msg = "🏦 <b>VAULT (ရန်ပုံငွေ အခြေအနေချုပ်)</b>\n━━━━━━━━━━━━━━\n"
    
    grand_total = 0
    for f in funds:
        f_id = f.get("fund_id")
        if not f_id: continue
        f_name = f.get("fund_name", f_id)
        total = sum(s.get("paid_funds", {}).get(f_id, 0) for s in students)
        grand_total += total
        msg += f"💰 <b>{f_name}</b>: {total:,} MMK\n"
        
    msg += f"\n📊 <b>Grand Total: {grand_total:,} MMK</b>"
    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)

# ================= 📊 ရန်ပုံငွေ အခြေအနေခွဲကြည့်ရန် (/report) =================
async def report_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_super_admin(update.effective_user.id): return
    if not context.args: return await update.message.reply_text("⚠️ အသုံးပြုနည်း: /report <FundID>")
    
    fund_id = context.args[0].lower()
    db = get_db()
    
    fund = await db.funds.find_one({"fund_id": fund_id})
    if not fund: return await update.message.reply_text("⚠️ ထိုအမည်ဖြင့် ရန်ပုံငွေ မရှိပါ။")
    
    target = fund.get("target_amount", 0)
    students = await db.student_info.find({}).to_list(length=None)
    
    paid_list, partial_list, unpaid_list = [], [], []
    total = 0
    
    for s in students:
        paid = s.get("paid_funds", {}).get(fund_id, 0)
        roll = s.get("roll_no", "Unknown")
        total += paid
        match = re.search(r'(\d+)$', roll)
        roll_num = int(match.group(1)) if match else 9999
        data = {"roll": roll, "roll_num": roll_num, "paid": paid}
        
        if paid >= target: paid_list.append(data)
        elif paid > 0: partial_list.append(data)
        else: unpaid_list.append(data)
            
    paid_list.sort(key=lambda x: x["roll_num"])
    partial_list.sort(key=lambda x: x["roll_num"])
    unpaid_list.sort(key=lambda x: x["roll_num"])
    
    p_str = ', '.join([x["roll"] for x in paid_list])
    pt_str = ', '.join([f"{x['roll']} ({x['paid']})" for x in partial_list])
    up_str = ', '.join([x["roll"] for x in unpaid_list])

    msg = f"📊 <b>{fund.get('fund_name', fund_id)} Report</b>\n━━━━━━━━━━━━━━\n"
    msg += f"🟢 <b>Paid ({len(paid_list)}):</b> {p_str if p_str else '-'}\n\n"
    msg += f"🟡 <b>Partial ({len(partial_list)}):</b> {pt_str if pt_str else '-'}\n\n"
    msg += f"🔴 <b>Unpaid ({len(unpaid_list)}):</b> {up_str if up_str else '-'}\n\n💰 <b>Total Collected:</b> {total:,} MMK"
    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)

# ================= 🔔 ကြွေးတောင်း သတိပေးစာပို့ရန် (/remind) =================
async def remind_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_super_admin(update.effective_user.id): return
    if not context.args: return await update.message.reply_text("⚠️ အသုံးပြုနည်း: /remind <FundID>")
    
    fund_id = context.args[0].lower()
    db = get_db()
    
    fund = await db.funds.find_one({"fund_id": fund_id})
    if not fund: return await update.message.reply_text("⚠️ ထိုအမည်ဖြင့် ရန်ပုံငွေ မရှိပါ။")
    
    target = fund.get("target_amount", 0)
    fund_name = fund.get("fund_name", fund_id)
    sent_count = 0
    
    students = await db.student_info.find({}).to_list(length=None)
    
    for s in students:
        paid = s.get("paid_funds", {}).get(fund_id, 0)
        user_id = s.get("user_id")
        
        if paid < target and user_id:
            text = f"🔔 <b>Class Fund Reminder</b>\n\nမင်္ဂလာပါ။ <b>{fund_name}</b> အတွက် ရန်ပုံငွေ ထည့်ဝင်ရန် ကျန်ရှိနေသေးသဖြင့် သတိပေးအပ်ပါသည်။\n\n💰 ပေးသွင်းရန်ကျန်ငွေ: <b>{target - paid:,} MMK</b>\n\n<i>(အတန်းကိုယ်စားလှယ်ထံသို့ အမြန်ဆုံး ဆက်သွယ်ပေးသွင်း주시ပါရန်)</i>"
            try:
                await context.bot.send_message(chat_id=user_id, text=text, parse_mode=ParseMode.HTML)
                sent_count += 1
            except Exception as e:
                logger.error(f"Failed to send reminder to {user_id}: {e}")
                
    await update.message.reply_text(f"✅ စုစုပေါင်း ကျောင်းသား ({sent_count}) ဦးထံသို့ သတိပေးစာ ပို့ပြီးပါပြီ。")

# ================= 📂 စာရင်းအပြည့်အစုံ CSV ထုတ်ရန် (/export) =================
async def export_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_super_admin(update.effective_user.id): return
    
    db = get_db()
    students = await db.student_info.find({}).to_list(length=None)
    if not students: return await update.message.reply_text("စာရင်းမရှိပါ။")
    
    # ရန်ပုံငွေ ခေါင်းစဉ်များကို ရယူခြင်း
    funds = await db.funds.find({}).to_list(length=None)
    fund_ids = [f.get("fund_id") for f in funds if f.get("fund_id")]
    fund_names = [f.get("fund_name", f_id) for f, f_id in zip(funds, fund_ids)]
    
    # ခုံနံပါတ်စဉ်အလိုက် စီခြင်း
    students.sort(key=lambda x: int(re.search(r'(\d+)$', x.get("roll_no", "0")).group(1)) if re.search(r'(\d+)$', x.get("roll_no", "0")) else 9999)
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    # 📌 Admin လိုချင်သည့်အတိုင်း Name, Roll No သာထားပြီး Fund Columns များကိုပါ ပေါင်းထည့်ပေးခြင်း
    headers = ["Roll No", "Name"] + fund_names
    writer.writerow(headers)
    
    for s in students:
        row = [s.get("roll_no", ""), s.get("name", "")]
        for f_id in fund_ids:
            paid = s.get("paid_funds", {}).get(f_id, 0)
            row.append(paid)
        writer.writerow(row)
        
    output.seek(0)
    await update.message.reply_document(
        document=output.getvalue().encode('utf-8-sig'), 
        filename="Class_Full_Report.csv", 
        caption="✅ <b>ကျောင်းသားစာရင်းနှင့် ရန်ပုံငွေများ အပြည့်အစုံ ထုတ်ပြီးပါပြီ။</b>\n<i>(Phone, Telegram, Hostel များ ဖြုတ်ထားပါသည်)</i>", 
        parse_mode=ParseMode.HTML
    )

# ================= 📂 သတ်မှတ် Fund ID အလိုက် Excel (.xlsx) ထုတ်ယူခြင်း (/exportfund) =================
async def export_fund_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_super_admin(update.effective_user.id): return
    if len(context.args) < 1: return await update.message.reply_text("⚠️ <b>အသုံးပြုနည်း:</b> <code>/exportfund &lt;FundID&gt;</code>", parse_mode=ParseMode.HTML)
        
    fund_id = context.args[0].lower()
    db = get_db()
    
    fund = await db.funds.find_one({"fund_id": fund_id})
    if not fund: return await update.message.reply_text(f"⚠️ <code>{fund_id}</code> အမည်ဖြင့် ရန်ပုံငွေ မရှိပါ။", parse_mode=ParseMode.HTML)
        
    fund_name = fund.get("fund_name", fund_id)
    students = await db.student_info.find({}).to_list(length=None)
    if not students: return await update.message.reply_text("⚠️ Database ထဲတွင် ကျောင်းသား အချက်အလက် မရှိသေးပါ။")

    export_data, total_collected = [], 0
    for s in students:
        roll = s.get("roll_no", "➖")
        match = re.search(r'(\d+)$', roll)
        paid = s.get("paid_funds", {}).get(fund_id, 0)
        export_data.append({"roll": roll, "roll_num": int(match.group(1)) if match else 9999, "name": s.get("name", "➖"), "status": "Paid" if paid > 0 else "Unpaid", "paid": paid})
        total_collected += paid
        
    export_data.sort(key=lambda x: (-x["paid"], x["roll_num"]))
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = fund_name[:31]
    
    ws.append(["Roll No", "Name", "Status", "Amount"])
    for col in range(1, 5): ws.cell(row=1, column=col).font = Font(bold=True)
        
    green_fill, green_font = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"), Font(color="006100")
    red_fill, red_font = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"), Font(color="9C0006")
    
    for row_idx, d in enumerate(export_data, start=2):
        ws.cell(row=row_idx, column=1, value=d["roll"])
        ws.cell(row=row_idx, column=2, value=d["name"])
        status_cell = ws.cell(row=row_idx, column=3, value=d["status"])
        
        if d["status"] == "Paid": 
            status_cell.fill, status_cell.font = green_fill, green_font
        else: 
            status_cell.fill, status_cell.font = red_fill, red_font
            
        ws.cell(row=row_idx, column=4, value=d["paid"])
        
    last_row = len(export_data) + 3
    ws.cell(row=last_row, column=3, value="TOTAL:").font = Font(bold=True)
    ws.cell(row=last_row, column=4, value=total_collected).font = Font(bold=True)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    await update.message.reply_document(
        document=output, 
        filename=f"{fund_name}_Report.xlsx", 
        caption=f"✅ <b>{fund_name} အတွက် စာရင်းထုတ်ပြီးပါပြီ။</b>\n\n🔹 စုစုပေါင်း ရငွေ: <b>{total_collected:,} MMK</b>", 
        parse_mode=ParseMode.HTML
    )