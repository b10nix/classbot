# handlers/finance/reports.py

import asyncio
import io
import csv
import re
import logging
import html
from datetime import datetime
from telegram import Update, InputFile, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import ContextTypes
from telegram.constants import ParseMode

from database.db import get_db
from utils.auth import is_super_admin
from utils.constants import MAJORS
from services.audit_service import log_audit
from services.major_service import get_user_major, is_major_authorized
from services.finance_service import (
    build_major_filter,
    build_fund_major_filter,
    extract_paid_amount,
    extract_roll_number,
    compute_vault_summary,
    build_excel_fund_report,
    get_major_semester_context,
    fund_is_current_semester,
)

logger = logging.getLogger(__name__)


# ================= 📋 ကျောင်းသားစာရင်းကြည့်ရန် (/list) =================
async def list_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    major = await get_user_major(user_id)
    if not await is_major_authorized(user_id, major):
        return

    db = get_db(major=major)
    students = await db.student_info.find(build_major_filter(major)).to_list(length=None)
    if not students:
        return await update.message.reply_text(f"📭 <b>[{major.upper()}]</b> တွင် ကျောင်းသားစာရင်း မရှိသေးပါ။", parse_mode=ParseMode.HTML)

    registered_nums = {
        extract_roll_number(s.get("roll_no", ""))
        for s in students
        if extract_roll_number(s.get("roll_no", "")) != 9999
    }
    max_roll = max(registered_nums) if registered_nums else 59
    if max_roll < 59:
        max_roll = 59

    missing_rolls = [i for i in range(1, max_roll + 1) if i not in registered_nums]
    missing_str = ", ".join([f"{i:02d}" for i in missing_rolls])

    msg = (
        f"📊 <b>[{major.upper()}] ကျောင်းသား စာရင်းသွင်းမှု အခြေအနေ</b>\n━━━━━━━━━━━━━━━━━━━━━\n\n"
        f"🔹 <b>စာရင်းသွင်းပြီးသူ:</b> {len(students)} ယောက်\n"
        f"🔹 <b>စာရင်းသွင်းရန် ကျန်ရှိသူ:</b> {len(missing_rolls)} ယောက်\n\n"
        "⏳ <b>စာရင်းသွင်းရန် ကျန်သေးသည့် ခုံနံပါတ်များ:</b>\n"
        f"<code>{missing_str if missing_str else 'အားလုံးသွင်းပြီးပါပြီ 🎉'}</code>"
    )
    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)


# ================= 🏠 အဆောင်နေ ကျောင်းသားများ (/hostel) =================
async def hostel_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    major = await get_user_major(user_id)
    if not await is_major_authorized(user_id, major):
        return

    db = get_db(major=major)
    query = {"$and": [build_major_filter(major), {"hostel": {"$in": ["Yes", "yes", True]}}]}
    students = await db.student_info.find(query).sort("roll_no", 1).to_list(length=None)
    if not students:
        return await update.message.reply_text(f"🏠 <b>[{major.upper()}]</b> တွင် အဆောင်နေ ကျောင်းသား မရှိသေးပါ။", parse_mode=ParseMode.HTML)

    msg = f"🏠 <b>[{major.upper()}] အဆောင်နေ ကျောင်းသားများ</b>\n━━━━━━━━━━━━━━\n"
    for s in students:
        msg += f"🔹 {s.get('roll_no')} - {s.get('name')}\n"
    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)


# ================= 🏦 ရန်ပုံငွေ အခြေအနေချုပ် (/vault) =================
async def vault_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    major = await get_user_major(user_id)
    if not major or not await is_major_authorized(user_id, major):
        return await update.message.reply_text("⛔ သက်ဆိုင်ရာ Major EC/Admin များသာ Finance Vault ကို ကြည့်ခွင့်ရှိသည်။")

    db = get_db(major=major)
    if db is None:
        return await update.message.reply_text("❌ Database connection error.")

    v = await compute_vault_summary(major)

    msg = (
        f"🏦 <b>[{v['major']}] VAULT (လက်ရှိ ရန်ပုံငွေ အခြေအနေ)</b>\n"
        f"━━━━━━━━━━━━━━\n"
        f"💰 <b>Carry-over Balance ({v['semester_label']}):</b> {v['carry_over_amount']:,.0f} MMK\n\n"
        f"📥 <b>စုစုပေါင်း ကောက်ခံရငွေ:</b> {v['total_collected']:,.0f} MMK\n"
        f"📤 <b>EC များ သုံးစွဲပြီးငွေ:</b> {v['total_spent']:,.0f} MMK\n"
        f"📊 <b>လက်ကျန်ငွေ (Net Balance):</b> {v['net_balance']:,.0f} MMK"
    )
    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)


# ================= 📊 ရန်ပုံငွေ အခြေအနေခွဲကြည့်ရန် (/report) =================
async def report_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    major = await get_user_major(user_id)
    if not await is_major_authorized(user_id, major):
        return

    if not context.args:
        return await update.message.reply_text(
            "⚠️ <b>အသုံးပြုနည်း:</b> <code>/report &lt;FundID&gt;</code>",
            parse_mode=ParseMode.HTML
        )

    fund_id = context.args[0].lower()
    db = get_db(major=major)
    m_filter = build_major_filter(major)

    fund = await db.funds.find_one({"$and": [build_fund_major_filter(major), {"fund_id": re.compile(f"^{re.escape(fund_id)}$", re.IGNORECASE)}]})
    if not fund:
        return await update.message.reply_text(f"⚠️ <b>[{major.upper()}]</b> တွင် ထိုအမည်ဖြင့် ရန်ပုံငွေ မရှိပါ။", parse_mode=ParseMode.HTML)

    target = fund.get("target_amount", 0)
    students = await db.student_info.find(m_filter).to_list(length=None)

    paid_list, partial_list, unpaid_list = [], [], []
    total = 0

    for s in students:
        paid_amount = extract_paid_amount(s.get("paid_funds", {}).get(fund_id, 0))
        roll = s.get("roll_no", "Unknown")
        total += paid_amount

        data = {"roll": roll, "roll_num": extract_roll_number(roll), "paid": paid_amount}

        if paid_amount >= target and target > 0:
            paid_list.append(data)
        elif paid_amount > 0:
            partial_list.append(data)
        else:
            unpaid_list.append(data)

    paid_list.sort(key=lambda x: x["roll_num"])
    partial_list.sort(key=lambda x: x["roll_num"])
    unpaid_list.sort(key=lambda x: x["roll_num"])

    p_str = ', '.join([x["roll"] for x in paid_list])
    pt_str = ', '.join([f"{x['roll']} ({x['paid']:,})" for x in partial_list])
    up_str = ', '.join([x["roll"] for x in unpaid_list])

    msg = f"📊 <b>[{major.upper()}] {fund.get('fund_name', fund_id.capitalize())} Report</b>\n━━━━━━━━━━━━━━\n"
    msg += f"🟢 <b>Paid ({len(paid_list)}):</b> {p_str if p_str else '-'}\n\n"
    msg += f"🟡 <b>Partial ({len(partial_list)}):</b> {pt_str if pt_str else '-'}\n\n"
    msg += f"🔴 <b>Unpaid ({len(unpaid_list)}):</b> {up_str if up_str else '-'}\n\n"
    msg += f"💰 <b>Total Collected:</b> {total:,} MMK"

    await update.message.reply_text(msg, parse_mode=ParseMode.HTML)


# ================= 🔔 ငွေမသွင်းရသေးသူများကို သတိပေးရန် (/remind) =================
async def remind_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    major = await get_user_major(user_id)
    if not await is_major_authorized(user_id, major):
        return

    if not context.args:
        return await update.message.reply_text(
            "⚠️ <b>အသုံးပြုနည်း:</b> <code>/remind &lt;FundID&gt;</code>\n(ဥပမာ - <code>/remind sports</code>)",
            parse_mode=ParseMode.HTML
        )

    fund_id = context.args[0].lower()
    db = get_db(major=major)
    m_filter = build_major_filter(major)

    fund = await db.funds.find_one({"$and": [build_fund_major_filter(major), {"fund_id": re.compile(f"^{re.escape(fund_id)}$", re.IGNORECASE)}]})
    if not fund:
        return await update.message.reply_text(
            f"⚠️ <b>[{major.upper()}]</b> တွင် <code>{fund_id}</code> အမည်ဖြင့် ရန်ပုံငွေ မရှိပါ။",
            parse_mode=ParseMode.HTML
        )

    target_amount = fund.get("target_amount", 0)
    fund_name = fund.get("fund_name", fund_id)

    raw_students = await db.student_info.find(m_filter).to_list(length=None)
    if not raw_students:
        return await update.message.reply_text(f"⚠️ <b>[{major.upper()}]</b> Database တွင် ကျောင်းသားစာရင်း မရှိသေးပါ။", parse_mode=ParseMode.HTML)

    unique_students = {}
    for st in raw_students:
        r_no = st.get("roll_no")
        if not r_no:
            continue
        existing = unique_students.get(r_no)
        has_id = bool(st.get("user_id") or st.get("tg_id"))
        existing_has_id = bool(existing.get("user_id") or existing.get("tg_id")) if existing else False
        if not existing or (not existing_has_id and has_id):
            unique_students[r_no] = st

    students = list(unique_students.values())
    await update.message.reply_text(f"⏳ <b>[{major.upper()}] {fund_name}</b> အတွက် သတိပေးစာ စတင်ပို့ဆောင်နေပါသည်...", parse_mode=ParseMode.HTML)

    success_count = 0
    failed_students = []

    for student in students:
        paid = extract_paid_amount(student.get("paid_funds", {}).get(fund_id, 0))
        remaining = target_amount - paid

        if remaining > 0:
            stu_uid = student.get("user_id") or student.get("tg_id")
            roll_no = student.get("roll_no", "Unknown")
            tg_uname = student.get("tg_username", "")
            display_fail = f"{roll_no} (@{tg_uname})" if tg_uname else f"{roll_no} (No Username)"

            if not stu_uid:
                failed_students.append(display_fail)
                continue

            remind_text = (
                f"🔔 <b>[{major.upper()}] FUND REMINDER</b>\n━━━━━━━━━━━━━━━━━━\n"
                f"👤 အမည်: <b>{student.get('name')}</b>\n"
                f"🏷️ ရန်ပုံငွေ: <b>{fund_name}</b>\n\n"
                f"🎯 သတ်မှတ်ထားသော ပမာဏ: {target_amount:,} MMK\n"
                f"📥 သွင်းပြီးထားသော ပမာဏ: {paid:,} MMK\n"
                f"⚠️ <b>ထပ်မံသွင်းရန် လိုအပ်ငွေ: {remaining:,} MMK</b>\n━━━━━━━━━━━━━━━━━━\n"
                f"<i>ကျေးဇူးပြု၍ သတ်မှတ်ရက်အတွင်း EC များထံအမြန်ဆုံး ပေးသွင်းပေးပါရန် မေတ္တာရပ်ခံအပ်ပါသည်။</i>"
            )

            try:
                await context.bot.send_message(chat_id=stu_uid, text=remind_text, parse_mode=ParseMode.HTML)
                success_count += 1
                await asyncio.sleep(0.5)
            except Exception as e:
                logger.error(f"Failed to remind user {stu_uid}: {e}")
                failed_students.append(display_fail)

    failed_count = len(failed_students)
    summary_msg = (
        f"✅ <b>[{major.upper()}] သတိပေးခြင်း ပြီးဆုံးပါပြီ။</b>\n\n"
        f"📩 အောင်မြင်စွာ ပို့ပြီးသူ: <b>{success_count}</b> ဦး\n"
        f"❌ ပို့၍မရသူ စုစုပေါင်း: <b>{failed_count}</b> ဦး"
    )
    if failed_count > 0:
        fail_list_str = "\n".join([f"• {fs}" for fs in failed_students])
        summary_msg += f"\n\n<b>[ ⚠️ ပို့၍မရသူများ စာရင်း ]</b>\n{fail_list_str}"

    await update.message.reply_text(summary_msg, parse_mode=ParseMode.HTML)


# ================= 📂 STUDENT MASTER EXPORT (.XLSX) =================
_EXPORT_PREFIX = "exportpick:"


def _hostel_truth(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value or "").strip().lower() in {"true", "yes", "y", "1", "နေသည်", "နေတယ်", "hostel"}


def _build_student_export(sheets: dict[str, list[dict]]) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    green = PatternFill(fill_type="solid", fgColor="E2F0D9")
    red = PatternFill(fill_type="solid", fgColor="FCE4D6")
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for sheet_name, students in sheets.items():
        ws = wb.create_sheet((sheet_name or "Major")[:31])
        ws.append(["roll", "name", "ph", "hostel"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        students = sorted(students, key=lambda st: extract_roll_number(st.get("roll_no", "")))
        for student in students:
            is_hostel = _hostel_truth(student.get("hostel"))
            row_values = [
                (f"{extract_roll_number(student):02d}" if extract_roll_number(student) != 9999 and extract_roll_number(student) < 100 else (str(extract_roll_number(student)) if extract_roll_number(student) != 9999 else str(student.get("roll_no", "") or ""))),
                str(student.get("name", "") or ""),
                str(student.get("ph", student.get("phone", "")) or ""),
                is_hostel,
            ]
            ws.append(row_values)
            fill = green if is_hostel else red
            for cell in ws[ws.max_row]:
                cell.fill = fill
            ws.cell(row=ws.max_row, column=4).number_format = '"TRUE";[Red]"FALSE"'
        widths = [14, 32, 18, 12]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def _send_student_export(message, context, major_choice: str, requested_by: int):
    if major_choice == "all" and not is_super_admin(requested_by):
        return await message.reply_text("⛔ All-Major export ကို Super Admin သာ သုံးနိုင်ပါတယ်။")
    if major_choice != "all" and major_choice not in MAJORS:
        return await message.reply_text("❌ Major မမှန်ပါ။")

    majors = MAJORS if major_choice == "all" else [major_choice]
    sheets: dict[str, list[dict]] = {}
    for major in majors:
        if not await is_major_authorized(requested_by, major):
            return await message.reply_text(f"⛔ [{major.upper()}] ကို Export လုပ်ခွင့်မရှိပါ။")
        db = get_db(major=major)
        if db is None:
            return await message.reply_text("❌ Database ချိတ်ဆက်မှု မအောင်မြင်ပါ။")
        students = await db.student_info.find(build_major_filter(major)).to_list(length=None)
        # Sheet names identify the Major, keeping the four data columns exactly as requested.
        sheets[major.upper()] = students
    total = sum(len(rows) for rows in sheets.values())
    if total == 0:
        return await message.reply_text("📭 Export လုပ်ရန် ကျောင်းသားစာရင်း မရှိပါ။")
    output = await asyncio.to_thread(_build_student_export, sheets)
    filename = f"{major_choice.upper()}_Students_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    caption = (f"✅ <b>{'All Majors' if major_choice == 'all' else major_choice.upper()} Student Export</b>\n"
               f"👥 ကျောင်းသား: <b>{total}</b> ဦး\n"
               "📄 Columns: <code>roll · name · ph · hostel</code>\n"
               "🟩 Hostel TRUE / 🟥 Hostel FALSE")
    db_for_audit = get_db()
    try:
        await log_audit(db_for_audit, major=major_choice if major_choice != "all" else "all",
                        actor_id=requested_by, action="student_export", entity_type="student_info",
                        entity_id=major_choice, details={"student_count": total, "majors": majors})
    except Exception:
        logger.exception("Could not log student export")
    await message.reply_document(document=InputFile(output, filename=filename), caption=caption, parse_mode=ParseMode.HTML)


async def export_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    if is_super_admin(user_id):
        if context.args:
            choice = context.args[0].lower()
            if choice == "all":
                return await _send_student_export(update.message, context, "all", user_id)
            major = normalize_export_major(choice)
            if not major:
                return await update.message.reply_text("အသုံးပြုပုံ: `/export` → inline ရွေးရန်၊ သို့မဟုတ် `/export all`၊ `/export ceit`", parse_mode=ParseMode.MARKDOWN)
            return await _send_student_export(update.message, context, major, user_id)
        rows = [[InlineKeyboardButton("🌐 All Majors", callback_data=f"{_EXPORT_PREFIX}all")]]
        for major in MAJORS:
            rows.append([InlineKeyboardButton(major.upper(), callback_data=f"{_EXPORT_PREFIX}{major}")])
        return await update.message.reply_text(
            "📤 <b>Student Export</b>\nExport အတွက် Major ရွေးပါ။ All Majors ကိုရွေးရင် Excel ထဲမှာ Major တစ်ခုစီ သီးခြား worksheet နဲ့ ထွက်ပါမယ်။",
            reply_markup=InlineKeyboardMarkup(rows), parse_mode=ParseMode.HTML)
    major = normalize_export_major(await get_user_major(user_id))
    if not major or not await is_major_authorized(user_id, major):
        return await update.message.reply_text("⛔ ကိုယ့် Major အတွက်သာ Export လုပ်ခွင့်ရှိပါတယ်။")
    return await _send_student_export(update.message, context, major, user_id)


def normalize_export_major(value):
    from services.major_service import normalize_major
    return normalize_major(str(value).strip()) if value is not None else None


async def export_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user_id = query.from_user.id
    choice = query.data[len(_EXPORT_PREFIX):]
    if not is_super_admin(user_id):
        return await query.edit_message_text("⛔ Export Major ရွေးချယ်ခွင့်ကို Super Admin သာ အသုံးပြုနိုင်ပါတယ်။")
    if choice != "all" and choice not in MAJORS:
        return await query.edit_message_text("Major မမှန်ပါ။ `/export` ကို ပြန်စပါ။")
    # Replace the picker message with a progress note, then send the workbook to the same chat.
    await query.edit_message_text("⏳ Excel export ပြင်ဆင်နေပါသည်။")
    return await _send_student_export(query.message, context, choice, user_id)


# ================= 📂 သတ်မှတ် Fund ID အလိုက် Excel (.xlsx) ထုတ်ယူခြင်း (/exportfund) =================
async def export_fund_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    major = await get_user_major(user_id)
    if not await is_major_authorized(user_id, major):
        return
    if len(context.args) < 1:
        return await update.message.reply_text("⚠️ <b>အသုံးပြုနည်း:</b> <code>/exportfund &lt;FundID&gt;</code>", parse_mode=ParseMode.HTML)

    fund_id = context.args[0].lower()
    db = get_db(major=major)
    m_filter = build_major_filter(major)

    fund = await db.funds.find_one({"$and": [build_fund_major_filter(major), {"fund_id": re.compile(f"^{re.escape(fund_id)}$", re.IGNORECASE)}]})
    if not fund:
        return await update.message.reply_text(f"⚠️ <b>[{major.upper()}]</b> တွင် <code>{fund_id}</code> အမည်ဖြင့် ရန်ပုံငွေ မရှိပါ။", parse_mode=ParseMode.HTML)

    fund_name = fund.get("fund_name", fund_id)
    students = await db.student_info.find(m_filter).to_list(length=None)
    if not students:
        return await update.message.reply_text(f"⚠️ <b>[{major.upper()}]</b> Database ထဲတွင် ကျောင်းသား အချက်အလက် မရှိသေးပါ။", parse_mode=ParseMode.HTML)

    export_data, total_collected = [], 0
    for s in students:
        roll = s.get("roll_no", "➖")
        paid_amount = extract_paid_amount(s.get("paid_funds", {}).get(fund_id, 0))
        export_data.append({
            "roll": roll,
            "roll_num": extract_roll_number(roll),
            "name": s.get("name", "➖"),
            "status": "Paid" if paid_amount > 0 else "Unpaid",
            "paid": paid_amount
        })
        total_collected += paid_amount

    export_data.sort(key=lambda x: (-x["paid"], x["roll_num"]))
    output = await asyncio.to_thread(build_excel_fund_report, fund_name, export_data, total_collected)

    await update.message.reply_document(
        document=output,
        filename=f"{major.upper()}_{fund_name}_Report.xlsx",
        caption=f"✅ <b>[{major.upper()}] {fund_name} အတွက် စာရင်းထုတ်ပြီးပါပြီ။</b>\n\n🔹 စုစုပေါင်း ရငွေ: <b>{total_collected:,} MMK</b>",
        parse_mode=ParseMode.HTML
    )

# ================= 📚 OLD SEMESTER FINANCE HISTORY =================
_SEMHIST_PREFIX = "semhist:"


async def semhistory_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    if is_super_admin(user_id):
        if context.args:
            choice = normalize_export_major(context.args[0])
            if not choice:
                return await update.message.reply_text("အသုံးပြုပုံ: `/semhistory` သို့မဟုတ် `/semhistory ceit`", parse_mode=ParseMode.MARKDOWN)
            return await _show_semester_history(update.message, choice, edit=False)
        rows = [[InlineKeyboardButton(m.upper(), callback_data=f"{_SEMHIST_PREFIX}{m}")] for m in MAJORS]
        return await update.message.reply_text("📚 အဟောင်း Semester finance history ကြည့်ရန် Major ရွေးပါ။", reply_markup=InlineKeyboardMarkup(rows))
    major = normalize_export_major(await get_user_major(user_id))
    if not major or not await is_major_authorized(user_id, major):
        return await update.message.reply_text("⛔ ကိုယ့် Major history ကိုသာ ကြည့်ခွင့်ရှိပါတယ်။")
    return await _show_semester_history(update.message, major, edit=False)


async def _show_semester_history(target, major: str, edit: bool):
    db = get_db(major=major)
    if db is None:
        return await (target.edit_message_text("❌ Database connection error.") if edit else target.reply_text("❌ Database connection error."))
    current_semester, current_year = await get_major_semester_context(major)
    funds = await db.funds.find({"$and": [build_fund_major_filter(major), {"type": {"$ne": "opening_balance"}}]}).sort("created_at", -1).to_list(length=None)
    old_funds = [f for f in funds if not fund_is_current_semester(f, current_semester, current_year)]
    if not old_funds:
        text = f"📭 [{major.upper()}] မှာ အရင် Semester finance history မရှိသေးပါ။"
        if edit: return await target.edit_message_text(text)
        return await target.reply_text(text)
    students = await db.student_info.find(build_major_filter(major)).to_list(length=None)
    lines = [f"📚 <b>[{major.upper()}] SEMESTER HISTORY</b>", "━━━━━━━━━━━━━━━━━━"]
    for fund in old_funds[:35]:
        fund_id = str(fund.get("fund_id") or "")
        fund_name = html.escape(str(fund.get("fund_name") or fund.get("title") or fund_id))
        amount = int(fund.get("target_amount") or 0)
        collected = sum(extract_paid_amount((st.get("paid_funds") or {}).get(fund_id, 0)) for st in students)
        semester = html.escape(str(fund.get("semester") or fund.get("closed_semester") or "legacy"))
        status = html.escape(str(fund.get("status", "legacy")))
        lines.append(f"• <b>{fund_name}</b> — {semester} · {status}\n  Fund ID: <code>{html.escape(fund_id)}</code>\n  တစ်ဦးချင်း: {amount:,} MMK · စုစုပေါင်းရငွေ: {collected:,} MMK")
    archives = await db.semester_archives.find({"major": major}).sort("closed_at", -1).to_list(length=5)
    for archive in archives:
        lines.append(f"\n📦 Close record: <b>{html.escape(str(archive.get('semester_name', 'semester')))}</b> — {html.escape(str(archive.get('closed_at', '')))}")
    text = "\n\n".join(lines)
    if len(text) > 3900:
        text = text[:3800] + "\n…\n<i>ကျန်စာရင်းကို အတိုချုံ့ထားသည်။</i>"
    if edit:
        return await target.edit_message_text(text, parse_mode=ParseMode.HTML, reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("⬅️ Back", callback_data=f"{_SEMHIST_PREFIX}back")]]))
    return await target.reply_text(text, parse_mode=ParseMode.HTML)


async def semhistory_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user_id = query.from_user.id
    choice = query.data[len(_SEMHIST_PREFIX):]
    if choice == "back":
        if is_super_admin(user_id):
            rows = [[InlineKeyboardButton(m.upper(), callback_data=f"{_SEMHIST_PREFIX}{m}")] for m in MAJORS]
            return await query.edit_message_text("📚 Major ရွေးပါ။", reply_markup=InlineKeyboardMarkup(rows))
        own_major = normalize_export_major(await get_user_major(user_id))
        if own_major and await is_major_authorized(user_id, own_major):
            return await query.edit_message_text("`/semhistory` ကို ပြန် run လုပ်ပါ။", parse_mode=ParseMode.MARKDOWN)
        return await query.edit_message_text("⛔ Permission မရှိပါ။")
    if choice not in MAJORS:
        return await query.edit_message_text("Major မမှန်ပါ။")
    if not is_super_admin(user_id):
        own_major = normalize_export_major(await get_user_major(user_id))
        if own_major != choice or not await is_major_authorized(user_id, choice):
            return await query.edit_message_text("⛔ ကိုယ့် Major history ကိုသာ ကြည့်ခွင့်ရှိပါတယ်။")
    return await _show_semester_history(query, choice, edit=True)
