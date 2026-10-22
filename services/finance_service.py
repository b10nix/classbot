# services/finance_service.py

import io
import re
import logging
import uuid
import hashlib
import json
from datetime import datetime, timezone, timedelta
import openpyxl
from openpyxl.styles import PatternFill, Font

from database.db import get_db, is_major_database_enabled
from services.major_service import (
    DEFAULT_MAJOR,
    normalize_major,
    resolve_ec_nick_by_id_in_major
)
from services.roll_number_service import (
    RollNumberError,
    CrossMajorAccessError,
    parse_roll_number,
    normalize_roll_input,
    get_academic_setting,
    MAJOR_ALIASES,
)

logger = logging.getLogger(__name__)


# ================= 🛠️ PURE DATA HELPERS =================
def extract_paid_amount(paid_data) -> int:
    if isinstance(paid_data, dict):
        return int(paid_data.get("amount", 0))
    elif isinstance(paid_data, (int, float)):
        return int(paid_data)
    return 0


def payment_entry_fingerprint(entry: dict) -> str:
    """Return a stable fingerprint for a payment-history entry, including legacy entries without IDs."""
    payload = json.dumps(entry, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def extract_roll_number(roll_str_or_doc: str | dict) -> int:
    """Extracts integer roll number from student dict or roll string."""
    if isinstance(roll_str_or_doc, dict):
        if isinstance(roll_str_or_doc.get("roll_number"), int):
            return roll_str_or_doc["roll_number"]
        roll_str_or_doc = roll_str_or_doc.get("roll_no", "")
    try:
        parsed = parse_roll_number(str(roll_str_or_doc))
        return parsed["roll_number"]
    except RollNumberError:
        match = re.search(r'(\d+)$', str(roll_str_or_doc))
        return int(match.group(1)) if match else 9999


def get_short_roll(roll_str_or_doc: str | dict) -> str:
    """Returns 2-digit padded roll string (e.g. '01', '43')."""
    if isinstance(roll_str_or_doc, dict):
        if roll_str_or_doc.get("roll_number_str"):
            return str(roll_str_or_doc["roll_number_str"])
        if isinstance(roll_str_or_doc.get("roll_number"), int):
            return f"{roll_str_or_doc['roll_number']:02d}"
        roll_str_or_doc = roll_str_or_doc.get("roll_no", "")
    try:
        parsed = parse_roll_number(str(roll_str_or_doc))
        return parsed["roll_str"]
    except RollNumberError:
        roll_str = str(roll_str_or_doc or "")
        return roll_str.split("-")[-1].strip() if "-" in roll_str else roll_str[-2:]


def build_major_filter(major: str) -> dict:
    """
    Builds a MongoDB filter for a major partition.
    For 'ceit', also matches legacy documents that don't have a 'major' field yet.
    """
    m = normalize_major(major) or DEFAULT_MAJOR
    # Legacy records without a major are accepted only in legacy single-DB mode.
    # Once separate major databases are enabled, untagged records must be fixed
    # by the reviewed migration, never silently treated as CEIT.
    if not is_major_database_enabled() and m == "ceit":
        return {"$or": [{"major": "ceit"}, {"major": {"$exists": False}}]}
    return {"major": m}


def build_fund_major_filter(major: str) -> dict:
    """Strictly scope fund documents to an explicitly assigned major.

    Unlike student records, a fund without a ``major`` value cannot safely be
    inferred to belong to CEIT in the shared database. Such legacy/unassigned
    fund documents are intentionally excluded from normal finance menus and
    operations until an admin reviews and assigns their owner.
    """
    m = normalize_major(major)
    if not m:
        raise ValueError("A valid major is required to query fund documents.")
    # Accept explicit legacy aliases and historical casing variants, but never
    # untagged or other-major funds.
    aliases = [alias for alias, canonical in MAJOR_ALIASES.items() if canonical == m]
    explicit_values: list[str] = []
    for alias in aliases:
        for value in (alias.lower(), alias.upper(), alias.capitalize()):
            if value not in explicit_values:
                explicit_values.append(value)
    return {"major": {"$in": explicit_values}}


# ================= 🗄️ SEMESTER-SCOPED FUND METADATA =================
async def get_major_semester_context(major: str = DEFAULT_MAJOR) -> tuple[str | int | None, str | None]:
    """Return (current_semester, academic_year) for a major, without guessing if unset."""
    m = normalize_major(major) or DEFAULT_MAJOR
    try:
        settings = await get_academic_setting(m)
    except Exception:
        logger.exception("Could not read academic setting for %s", m)
        settings = None
    if not settings:
        return None, None
    return settings.get("current_semester"), settings.get("academic_year")


def _same_semester_value(left, right) -> bool:
    if left is None or right is None:
        return left is right
    return str(left).strip().lower().replace("sem-", "") == str(right).strip().lower().replace("sem-", "")


def fund_is_current_semester(fund: dict, current_semester=None, current_year=None) -> bool:
    """Exclude explicitly closed funds and funds belonging to another configured semester."""
    status = str(fund.get("status", "active")).lower()
    if status in {"completed", "cancelled", "archived", "closed", "deleted"}:
        return False
    fund_semester = fund.get("semester")
    fund_year = fund.get("academic_year")
    # Legacy active funds without semester metadata remain visible until explicitly closed.
    if current_semester is None or fund_semester is None:
        return status == "active"
    if not _same_semester_value(fund_semester, current_semester):
        return False
    if current_year and fund_year and str(fund_year).strip() != str(current_year).strip():
        return False
    return True


async def get_funds_metadata(major: str = DEFAULT_MAJOR) -> tuple[list[dict], set[str], dict[str, str], dict[str, str]]:
    """Return current-semester funds plus a set of non-current/closed IDs.

    Historic fund names/owners are retained for historical reports, but default
    dashboard calculations only count current-semester, active fund IDs.
    """
    m = normalize_major(major) or DEFAULT_MAJOR
    db = get_db(major=m)
    if db is None:
        return [], set(), {}, {}
    fund_filter = build_fund_major_filter(m)
    query = {"$and": [fund_filter, {"type": {"$ne": "opening_balance"}}]}
    all_funds = await db.funds.find(query).to_list(length=None)
    current_semester, current_year = await get_major_semester_context(m)
    current_funds = []
    inactive_or_historic_ids: set[str] = set()
    fund_owners: dict[str, str] = {}
    fund_name_map: dict[str, str] = {}
    sys_config = await db.sys_config.find_one({"type": "archived_funds"}) if m == "ceit" or is_major_database_enabled() else None
    legacy_archived = {str(x).lower() for x in (sys_config or {}).get("funds", [])}

    for fund in all_funds:
        fund_id = str(fund.get("fund_id") or "").lower()
        if not fund_id:
            continue
        fund_name_map[fund_id] = fund.get("fund_name", fund.get("title", fund_id))
        manager_id = fund.get("manager_user_id")
        manager_name = fund.get("managed_by")
        if fund.get("assignment_mode") == "single_manager":
            manager_key = str(manager_id) if manager_id else (str(manager_name).lower() if manager_name else "ec")
        else:
            manager_key = "ec"  # monthly/assigned funds follow each student's allocation
        fund_owners[fund_id] = manager_key
        if fund_id in legacy_archived or not fund_is_current_semester(fund, current_semester, current_year):
            inactive_or_historic_ids.add(fund_id)
        else:
            current_funds.append(fund)
    return current_funds, inactive_or_historic_ids, fund_owners, fund_name_map


# ================= 📊 UNIFIED EC DASHBOARD ENGINE (MAJOR-SCOPED) =================
async def compute_ec_dashboard(
    target_user_id: int | None,
    target_identifiers: list[str],
    major: str = DEFAULT_MAJOR
) -> dict:
    m = normalize_major(major) or DEFAULT_MAJOR
    db = get_db(major=m)
    if db is None:
        return {"major": m, "has_records": False, "assigned_students_count": 0,
                "active_fund_totals": {}, "archived_fund_totals": {}, "fund_name_map": {},
                "grand_total_collected": 0, "total_spent": 0, "transferred_in": 0,
                "transferred_out": 0, "opening_note": None, "total_income": 0,
                "total_outflow": 0, "net_in_hand": 0}

    clean_identifiers = {str(x).lower().lstrip("@") for x in target_identifiers if x}
    if target_user_id:
        clean_identifiers.add(str(target_user_id))
        nick = await resolve_ec_nick_by_id_in_major(target_user_id, m)
        if nick != "ec":
            clean_identifiers.add(nick)

    current_funds, archived_fund_ids, fund_owners, fund_name_map = await get_funds_metadata(m)
    current_fund_ids = {str(f.get("fund_id") or "").lower() for f in current_funds}
    all_students = await db.student_info.find(build_major_filter(m)).to_list(length=None)

    assigned_students_count = 0
    active_fund_totals = {}
    archived_fund_totals = {}
    grand_total_collected = 0

    for st in all_students:
        is_target_student = False
        st_ec = str(st.get("assigned_ec", "")).lower()
        st_ec_id = st.get("assigned_ec_id")

        if st_ec in clean_identifiers or str(st_ec_id).lower() in clean_identifiers:
            is_target_student = True
        elif target_user_id and (st_ec_id == target_user_id or st.get("tg_id") == target_user_id):
            is_target_student = True

        if is_target_student:
            assigned_students_count += 1

        paid_funds = st.get("paid_funds", {})
        for f_id, f_data in paid_funds.items():
            amount = extract_paid_amount(f_data)
            if amount <= 0:
                continue

            clean_f_id = str(f_id).lower()
            fund_manager = fund_owners.get(clean_f_id, "ec")

            is_target_money = (
                (fund_manager == "ec" and is_target_student)
                or (fund_manager in clean_identifiers)
            )

            # Old-semester and closed funds stay in student history but are hidden
            # from current /me and /ecstat totals.
            if is_target_money and clean_f_id in current_fund_ids:
                active_fund_totals[clean_f_id] = active_fund_totals.get(clean_f_id, 0) + amount
                grand_total_collected += amount

    id_or_queries = [{"username": {"$in": list(clean_identifiers)}}]
    if target_user_id:
        id_or_queries.extend([
            {"ec_id": target_user_id},
            {"ec_id": str(target_user_id)},
            {"user_id": target_user_id}
        ])

    ledger_query = {"$and": [build_major_filter(m), {"$or": id_or_queries}]}
    ec_ledger = await db.ec_ledgers.find_one(ledger_query) or {}

    total_spent = ec_ledger.get("total_spent", 0)
    transferred_in = ec_ledger.get("transferred_in", 0)
    transferred_out = ec_ledger.get("transferred_out", 0)

    total_spent = int(total_spent) if isinstance(total_spent, (int, float)) else 0
    transferred_in = int(transferred_in) if isinstance(transferred_in, (int, float)) else 0
    transferred_out = int(transferred_out) if isinstance(transferred_out, (int, float)) else 0

    total_income = grand_total_collected + transferred_in
    total_outflow = total_spent + transferred_out
    net_in_hand = total_income - total_outflow

    has_records = not (
        assigned_students_count == 0
        and total_income == 0
        and total_outflow == 0
        and not archived_fund_totals
    )

    return {
        "major": m,
        "has_records": has_records,
        "assigned_students_count": assigned_students_count,
        "active_fund_totals": active_fund_totals,
        "archived_fund_totals": archived_fund_totals,
        "fund_name_map": fund_name_map,
        "grand_total_collected": grand_total_collected,
        "total_spent": total_spent,
        "transferred_in": transferred_in,
        "transferred_out": transferred_out,
        "opening_note": ec_ledger.get("opening_note"),
        "total_income": total_income,
        "total_outflow": total_outflow,
        "net_in_hand": net_in_hand,
    }


def format_ec_dashboard_message(title_name: str, data: dict, viewed_by_admin: bool = False) -> str:
    admin_tag = " (Viewed by Admin)" if viewed_by_admin else ""
    major_tag = data.get("major", DEFAULT_MAJOR).upper()
    fund_name_map = data["fund_name_map"]

    msg = f"📊 <b>[{major_tag}] EC DASHBOARD - @{title_name}{admin_tag}</b>\n"
    msg += "━━━━━━━━━━━━━━━━━━\n"
    msg += f"👥 တာဝန်ယူထားသော ကျောင်းသား: <b>{data['assigned_students_count']}</b> ဦး\n\n"

    msg += "<b>[ 🎯 ရငွေစာရင်း (Income) ]</b>\n"
    if data["active_fund_totals"]:
        for f_id, f_amt in data["active_fund_totals"].items():
            display_name = fund_name_map.get(f_id, f_id)
            msg += f"• {display_name}: <b>{f_amt:,} MMK</b>\n"
    else:
        msg += "<i>လက်ရှိ ကောက်ခံထားသော ငွေမရှိသေးပါ</i>\n"

    if data["transferred_in"] > 0:
        in_label = data.get("opening_note") or "အခြား EC များထံမှ လက်ခံရရှိငွေ"
        msg += f"• {in_label}: <b>{data['transferred_in']:,} MMK</b>\n"

    msg += f"📥 <b>စုစုပေါင်း ရငွေ: {data['total_income']:,} MMK</b>\n\n"

    msg += "<b>[ 💸 သုံးစွဲမှုနှင့် ထွက်ငွေ (Outflow) ]</b>\n"
    msg += f"• သုံးစွဲပြီးငွေ (Spent): <b>{data['total_spent']:,} MMK</b>\n"

    if data["transferred_out"] > 0:
        msg += f"• အခြား EC များသို့ လွှဲပြောင်းပေးငွေ: <b>{data['transferred_out']:,} MMK</b>\n"

    msg += f"📤 <b>စုစုပေါင်း ထွက်ငွေ: {data['total_outflow']:,} MMK</b>\n\n"
    msg += f"💰 <b>လက်ထဲတွင် ကျန်ငွေ (Net Balance): {data['net_in_hand']:,} MMK</b>\n"

    if data["archived_fund_totals"]:
        msg += "\n📦 <b>[ ဆရာမထံ အပ်နှံပြီးသော ရန်ပုံငွေများ ]</b>\n"
        for f_id, f_amt in data["archived_fund_totals"].items():
            display_name = fund_name_map.get(f_id, f_id)
            msg += f"• {display_name}: <b>{f_amt:,} MMK</b>\n"

    return msg


# ================= 💳 STUDENT LOOKUP, PAYMENT & UNDO (MAJOR-SCOPED) =================
async def find_student_by_roll_in_major(
    roll_query: str,
    major: str = DEFAULT_MAJOR,
    allow_cross_major: bool = False
) -> dict | None:
    """
    Finds a student by numeric roll ('01', '1') or legacy/canonical roll ('Sem-III-CEIT-01').
    Enforces strict major isolation:
    - If `roll_query` contains an explicit major prefix (e.g. 'Sem-III-EP-01') that differs from
      `major` and `allow_cross_major` is False, raises CrossMajorAccessError.
    - Matches by integer `roll_number`, current `roll_no`, or historical `roll_history.roll_no`.
    """
    target_major = normalize_major(major) or DEFAULT_MAJOR
    parsed = parse_roll_number(roll_query)
    if parsed["major"] and parsed["major"] != target_major:
        if not allow_cross_major:
            raise CrossMajorAccessError(
                f"⛔ သင်သည် [{target_major.upper()}] အတွက်သာ လုပ်ပိုင်ခွင့်ရှိပြီး "
                f"[{parsed['major'].upper()}] ကျောင်းသားစာရင်းကို ပြင်ဆင်ခွင့် မရှိပါ။"
            )
        target_major = parsed["major"]

    db = get_db(major=target_major)
    if db is None:
        return None
    r_int = parsed["roll_number"]
    m_filter = build_major_filter(target_major)

    # If a full canonical/legacy string with a specific semester was passed, match either exact or roll_number
    if not parsed["is_numeric_only"] and parsed["canonical"]:
        exact_regex = re.compile(f"^{re.escape(str(roll_query).strip())}$", re.IGNORECASE)
        canon_regex = re.compile(f"^{re.escape(parsed['canonical'])}$", re.IGNORECASE)
        query = {
            "$and": [
                m_filter,
                {"$or": [
                    {"roll_no": canon_regex},
                    {"roll_no": exact_regex},
                    {"roll_history.roll_no": exact_regex},
                    {"roll_number": r_int}
                ]}
            ]
        }
    else:
        trailing_regex = re.compile(rf"(?:^|-)0*{r_int}$", re.IGNORECASE)
        query = {
            "$and": [
                m_filter,
                {"$or": [
                    {"roll_number": r_int},
                    {"roll_no": trailing_regex}
                ]}
            ]
        }

    return await db.student_info.find_one(query)


async def record_student_payment(
    roll_query: str,
    fund_id: str,
    amount: int,
    admin_id: int,
    admin_name: str,
    major: str = DEFAULT_MAJOR,
    allow_cross_major: bool = False,
    responsible_ec_id: int | str | None = None,
    responsible_ec_name: str | None = None,
    fund_semester: str | int | None = None,
) -> tuple[bool, str, dict | None]:
    m = normalize_major(major) or DEFAULT_MAJOR
    fund_id = str(fund_id or "").strip().casefold()
    db = get_db(major=m)
    if db is None:
        return False, "❌ Database connection is not available.", None

    try:
        student = await find_student_by_roll_in_major(roll_query, m, allow_cross_major=allow_cross_major)
    except CrossMajorAccessError as e:
        return False, str(e), None
    except RollNumberError as e:
        return False, f"⚠️ {e}", None

    if not student:
        return False, f"⚠️ <b>[{m.upper()}]</b> မေဂျာတွင် Roll <code>{roll_query}</code> ကျောင်းသားကို ရှာမတွေ့ပါ။", None

    student_major = normalize_major(student.get("major")) or m
    student_db = get_db(major=student_major)
    if student_db is None:
        return False, "❌ Database connection is not available.", None
    current_fund_data = student.get("paid_funds", {}).get(fund_id, 0)
    current_paid = 0
    collector_breakdown = {}

    if isinstance(current_fund_data, dict):
        current_paid = current_fund_data.get("amount", 0)
        old_collected_by = current_fund_data.get("collected_by", "ec")
        if isinstance(old_collected_by, str):
            collector_breakdown[old_collected_by] = current_paid
        elif isinstance(old_collected_by, dict):
            collector_breakdown = old_collected_by.copy()
    elif isinstance(current_fund_data, (int, float)):
        current_paid = int(current_fund_data)
        if current_paid > 0:
            collector_breakdown[student.get("assigned_ec", "ec")] = current_paid

    # A fund payment is a one-time collection per student/fund in this workflow.
    # Reject known duplicates, then enforce the same rule atomically in MongoDB
    # so two ECs confirming at the same time cannot write duplicate payments.
    if current_paid > 0:
        return False, f"⚠️ Roll <code>{get_short_roll(student)}</code> သည် ဒီ Fund ကို ပေးပြီးသားဖြစ်သည်။", None
    new_paid = amount
    admin_nick = await resolve_ec_nick_by_id_in_major(admin_id, student_major)
    collector_breakdown[admin_nick] = collector_breakdown.get(admin_nick, 0) + amount

    admin_info = f"{admin_name} ({admin_id})"
    mm_tz = timezone(timedelta(hours=6, minutes=30))
    current_dt = datetime.now(mm_tz)
    current_date = current_dt.strftime("%d-%b-%Y %I:%M %p")
    transaction_id = uuid.uuid4().hex

    paid_path = f"paid_funds.{fund_id}"
    update_result = await student_db.student_info.update_one(
        {"$and": [
            {"_id": student["_id"]},
            {"$or": [
                {paid_path: {"$exists": False}},
                {paid_path: 0},
                {f"{paid_path}.amount": 0},
            ]},
        ]},
        {
            "$set": {
                f"paid_funds.{fund_id}": {
                    "amount": new_paid,
                    "collected_by": collector_breakdown,
                    "last_date": current_date
                },
                "last_payment_date": current_date
            },
            "$push": {
                "payment_history": {
                    "transaction_id": transaction_id,
                    "fund_id": fund_id,
                    "amount": amount,
                    "date": current_date,
                    "received_by": admin_info,
                    "collector_nick": admin_nick,
                    "semester": fund_semester if fund_semester is not None else student.get("current_semester"),
                    "received_by_id": admin_id,
                    "responsible_ec_id": responsible_ec_id,
                    "responsible_ec_name": responsible_ec_name,
                    "roll_no_snapshot": get_short_roll(student)
                }
            }
        }
    )

    if getattr(update_result, "modified_count", 0) != 1:
        return False, f"⚠️ Roll <code>{get_short_roll(student)}</code> ကို အခြား EC တစ်ယောက်က အရင်မှတ်တမ်းတင်ပြီးဖြစ်နိုင်သည်။ ပြန်စစ်ပါ။", None

    # Keep a first-class transaction record for finance reporting while retaining
    # the existing student payment fields for backward-compatible dashboards.
    try:
        await student_db.fund_transactions.insert_one({
            "transaction_id": transaction_id,
            "type": "student_payment",
            "status": "recorded",
            "major": student_major,
            "semester": fund_semester if fund_semester is not None else student.get("current_semester"),
            "fund_id": {"$regex": f"^{re.escape(fund_id)}$", "$options": "i"},
            "student_ref": student["_id"],
            "internal_id": student.get("internal_id"),
            "roll_no": get_short_roll(student),
            "student_name": student.get("name"),
            "amount": amount,
            "received_by_id": admin_id,
            "received_by_name": admin_name,
            "collector_nick": admin_nick,
            "responsible_ec_id": responsible_ec_id,
            "responsible_ec_name": responsible_ec_name,
            "date": current_date,
            "created_at": current_dt.isoformat(),
        })
    except Exception:
        # The primary student payment has committed. Do not falsely report it as
        # failed; retain the embedded history and log the transaction-ledger issue.
        logger.exception(
            "Payment saved on student record but transaction ledger insert failed "
            "(major=%s fund=%s transaction_id=%s)",
            student_major, fund_id, transaction_id,
        )

    return True, "Success", {
        "internal_id": student.get("internal_id"),
        "name": student.get("name"),
        "roll_no": student.get("roll_no"),
        "major": student_major.upper(),
        "amount": amount,
        "new_paid": new_paid,
        "date": current_date
    }


async def undo_student_payment(
    roll_query: str,
    fund_id: str,
    major: str = DEFAULT_MAJOR,
    allow_cross_major: bool = False,
    reversed_by_id: int | None = None,
    expected_transaction_id: str | None = None,
    expected_transaction_fingerprint: str | None = None,
    expected_date: str | None = None,
    expected_amount: int | None = None,
) -> tuple[bool, str, dict | None]:
    m = normalize_major(major) or DEFAULT_MAJOR
    fund_id = str(fund_id or "").strip().casefold()
    db = get_db(major=m)
    if db is None:
        return False, "❌ Database connection is not available.", None

    try:
        student = await find_student_by_roll_in_major(roll_query, m, allow_cross_major=allow_cross_major)
    except CrossMajorAccessError as e:
        return False, str(e), None
    except RollNumberError as e:
        return False, f"⚠️ {e}", None

    if not student:
        return False, f"⚠️ <b>[{m.upper()}]</b> မေဂျာတွင် Roll <code>{roll_query}</code> ကျောင်းသားကို ရှာမတွေ့ပါ။", None

    student_major = normalize_major(student.get("major")) or m
    student_db = get_db(major=student_major)
    if student_db is None:
        return False, "❌ Database connection is not available.", None
    stored_history = student.get("payment_history", [])
    # Keep an untouched snapshot for compare-and-set; mutate a separate copy below.
    history = list(stored_history) if isinstance(stored_history, list) else []
    fund_history = [
        h for h in history
        if isinstance(h, dict) and str(h.get("fund_id") or "").casefold() == fund_id
    ]

    if not fund_history:
        return False, f"⚠️ <b>{fund_id}</b> အတွက် ပြန်ရုပ်သိမ်းစရာ ငွေသွင်းမှတ်တမ်း မရှိပါ။", None

    last_tx = fund_history[-1]
    # Inline confirmation passes the transaction shown to the EC. Recheck that
    # the same transaction is still the newest eligible payment for this fund.
    # This prevents a stale confirmation from undoing a newer payment.
    if expected_transaction_id and str(last_tx.get("transaction_id") or "") != str(expected_transaction_id):
        return False, "⚠️ အတည်ပြုထားတဲ့ payment မဟုတ်တော့ပါ။ /undo ကို ပြန်စပြီး အချက်အလက်အသစ်ကို စစ်ပါ။", None
    if expected_transaction_fingerprint and payment_entry_fingerprint(last_tx) != expected_transaction_fingerprint:
        return False, "⚠️ အတည်ပြုထားတဲ့ payment မှတ်တမ်း ပြောင်းလဲသွားပါတယ်။ /undo ကို ပြန်စပါ။", None
    if expected_date is not None and str(last_tx.get("date") or "") != str(expected_date):
        return False, "⚠️ Payment ရက်စွဲ ပြောင်းလဲသွားပါတယ်။ /undo ကို ပြန်စပြီး စစ်ပါ။", None
    if expected_amount is not None:
        try:
            if int(last_tx.get("amount") or 0) != int(expected_amount):
                return False, "⚠️ Payment ပမာဏ ပြောင်းလဲသွားပါတယ်။ /undo ကို ပြန်စပြီး စစ်ပါ။", None
        except (TypeError, ValueError):
            return False, "⚠️ Payment ပမာဏကို ပြန်စစ်လို့မရပါ။ /undo ကို ပြန်စပါ။", None
    revert_amount = last_tx["amount"]
    tx_date = last_tx["date"]

    collector_nick = last_tx.get("collector_nick")
    if not collector_nick:
        received_by_str = last_tx.get("received_by", "")
        id_match = re.search(r'\((\d+)\)', received_by_str)
        if id_match:
            collector_nick = await resolve_ec_nick_by_id_in_major(int(id_match.group(1)), student_major)
        else:
            collector_nick = "ec"

    stored_paid_funds = student.get("paid_funds") if isinstance(student.get("paid_funds"), dict) else {}
    stored_paid_key = next(
        (str(key) for key in stored_paid_funds if str(key).casefold() == fund_id),
        fund_id,
    )
    current_fund_data = stored_paid_funds.get(stored_paid_key, 0)
    current_paid = extract_paid_amount(current_fund_data)
    new_paid = max(0, current_paid - revert_amount)

    if isinstance(current_fund_data, dict):
        set_value = current_fund_data.copy()
        set_value["amount"] = new_paid
        collected_by = set_value.get("collected_by")
        if isinstance(collected_by, dict):
            updated_cb = collected_by.copy()
            if collector_nick in updated_cb:
                updated_cb[collector_nick] = max(0, updated_cb[collector_nick] - revert_amount)
                if updated_cb[collector_nick] == 0:
                    del updated_cb[collector_nick]
            elif updated_cb:
                first_k = next(iter(updated_cb))
                updated_cb[first_k] = max(0, updated_cb[first_k] - revert_amount)
            set_value["collected_by"] = updated_cb
    else:
        set_value = new_paid

    for i in range(len(history) - 1, -1, -1):
        item = history[i]
        if (
            isinstance(item, dict)
            and str(item.get("fund_id") or "").casefold() == fund_id
            and item.get("amount") == revert_amount
            and item.get("date") == tx_date
        ):
            history.pop(i)
            break

    # Compare-and-set against the exact payment snapshot read above. Without this
    # guard, two ECs undoing at nearly the same time could overwrite each other's
    # payment_history / paid_funds updates because the old query only matched _id.
    undo_filter = {"_id": student["_id"], "payment_history": stored_history}
    if "paid_funds" in student:
        undo_filter["paid_funds"] = student.get("paid_funds")
    else:
        undo_filter["paid_funds"] = {"$exists": False}

    student_update = await student_db.student_info.update_one(
        undo_filter,
        {
            "$set": {
                f"paid_funds.{stored_paid_key}": set_value,
                "payment_history": history
            }
        }
    )
    if getattr(student_update, "modified_count", 0) != 1:
        return False, (
            "⚠️ ဒီ payment မှတ်တမ်းကို အခြား EC တစ်ဦးက အရင်ပြောင်းလဲထားနိုင်ပါတယ်။ "
            "အချက်အလက်ကို ပြန်စစ်ပြီးမှ ထပ်လုပ်ပါ။"
        ), None

    # Reverse the corresponding first-class transaction instead of deleting it.
    try:
        transaction_filter = {
            "major": student_major,
            "type": "student_payment",
            "fund_id": {"$regex": f"^{re.escape(fund_id)}$", "$options": "i"},
            "student_ref": student["_id"],
            "status": "recorded",
        }
        if last_tx.get("transaction_id"):
            transaction_filter["transaction_id"] = last_tx["transaction_id"]
        else:
            transaction_filter["date"] = tx_date
        stored_transaction = await student_db.fund_transactions.find_one(
            transaction_filter, sort=[("created_at", -1)]
        )
        if stored_transaction:
            mm_tz = timezone(timedelta(hours=6, minutes=30))
            await student_db.fund_transactions.update_one(
                {"_id": stored_transaction["_id"], "status": "recorded"},
                {"$set": {
                    "status": "reversed",
                    "reversed_at": datetime.now(mm_tz).isoformat(),
                    "reversed_by_id": reversed_by_id,
                    "reversed_amount": revert_amount,
                }},
            )
    except Exception:
        logger.exception(
            "Student payment reversed but transaction status update failed "
            "(major=%s fund=%s roll=%s)", student_major, fund_id, get_short_roll(student),
        )

    return True, "Success", {
        "name": student.get("name"),
        "roll_no": student.get("roll_no"),
        "major": student_major.upper(),
        "revert_amount": revert_amount,
        "tx_date": tx_date,
        "new_paid": new_paid
    }


# ================= 🏦 VAULT & REPORTS COMPUTATION (MAJOR-SCOPED) =================
async def compute_vault_summary(major: str = DEFAULT_MAJOR) -> dict:
    m = normalize_major(major) or DEFAULT_MAJOR
    db = get_db(major=m)
    if db is None:
        return {"major": m.upper(), "semester_label": "Current", "carry_over_amount": 0,
                "total_collected": 0, "total_spent": 0, "net_balance": 0}
    m_filter = build_major_filter(m)
    current_semester, current_year = await get_major_semester_context(m)

    opening_query = {"$and": [m_filter, {"$or": [{"type": "opening_balance"}, {"is_opening": True}]}]}
    if current_semester is not None:
        opening_query["$and"].append({"$or": [{"semester": current_semester}, {"semester": {"$exists": False}}]})
    opening_tx = await db.fund_transactions.find(opening_query).sort("date", -1).to_list(length=1)
    opening_tx = opening_tx[0] if opening_tx else None

    carry_over_amount = 0.0
    semester_label = str(current_semester) if current_semester is not None else "Current"
    if opening_tx:
        carry_over_amount = float(opening_tx.get("amount", 0.0) or 0.0)
        desc = opening_tx.get("description", "")
        if "from " in desc:
            semester_label = f"{semester_label} (carry-over from {desc.split('from ')[-1].strip()})"
    else:
        opening_query = {"$and": [build_fund_major_filter(m), {"type": "opening_balance"}]}
        if current_semester is not None:
            opening_query["$and"].append({"$or": [{"semester": current_semester}, {"semester": {"$exists": False}}]})
        opening_fund = await db.funds.find_one(opening_query)
        if opening_fund:
            carry_over_amount = float(opening_fund.get("amount") or opening_fund.get("target_amount") or 0.0)

    current_funds, _, _, _ = await get_funds_metadata(m)
    current_ids = {str(f.get("fund_id") or "").lower() for f in current_funds}
    students = await db.student_info.find(m_filter).to_list(length=None)
    total_collected = 0.0
    for student in students:
        for fund_id, fund_data in (student.get("paid_funds") or {}).items():
            if str(fund_id).lower() in current_ids:
                total_collected += float(extract_paid_amount(fund_data))

    ec_ledgers = await db.ec_ledgers.find(m_filter).to_list(length=None)
    total_spent = sum(
        float(ec.get("total_spent", 0))
        for ec in ec_ledgers
        if isinstance(ec.get("total_spent", 0), (int, float))
    )
    net_balance = (carry_over_amount + total_collected) - total_spent
    return {
        "major": m.upper(), "semester_label": semester_label,
        "carry_over_amount": carry_over_amount, "total_collected": total_collected,
        "total_spent": total_spent, "net_balance": net_balance,
    }


def build_excel_fund_report(fund_name: str, export_data: list[dict], total_collected: int) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = fund_name[:31]

    ws.append(["Roll No", "Name", "Status", "Amount"])
    for col in range(1, 5):
        ws.cell(row=1, column=col).font = Font(bold=True)

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(color="006100")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006")

    for row_idx, d in enumerate(export_data, start=2):
        ws.cell(row=row_idx, column=1, value=d["roll"])
        ws.cell(row=row_idx, column=2, value=d["name"])
        status_cell = ws.cell(row=row_idx, column=3, value=d["status"])

        if d["status"] == "Paid":
            status_cell.fill, status_cell.font = green_fill, green_font
        else:
            status_cell.fill, status_cell.font = red_fill, red_font

        ws.cell(row=row_idx, column=4, value=d["paid"])

    last_row = len(export_data) + 3
    ws.cell(row=last_row, column=3, value="TOTAL:").font = Font(bold=True)
    ws.cell(row=last_row, column=4, value=total_collected).font = Font(bold=True)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output